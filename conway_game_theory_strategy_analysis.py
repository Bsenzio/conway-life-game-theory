#!/usr/bin/env python3
"""Phase 4: reproducible two-player game-theoretic experiments on Conway B3/S23.

Install: py -m pip install numpy scipy matplotlib openpyxl
Run:     py conway_game_theory_phase4.py
Optional: py conway_game_theory_phase4.py --output MyResults

IMPORTANT: These are deliberately finite 3x3 *restricted* strategy games, not
proofs of equilibria over every possible 5x5 cell configuration.
"""
import argparse
import csv
import itertools
import json
import math
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
from scipy.optimize import linprog
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SIZE = 5
OBSERVED_GENERATIONS = 10  # Generation 1 = initial, Generation 10 = nine transitions.
BUDGETS = (3, 4, 5, 6)
FAMILIES = ('Compact', 'Row-oriented', 'Dispersed')
TOL = 1e-7


def evolve(cells):
    board = frozenset(cells)
    states = [board]
    for _ in range(OBSERVED_GENERATIONS - 1):
        nxt = set()
        for r in range(SIZE):
            for c in range(SIZE):
                neighbors = sum((rr, cc) in board
                                for rr in range(max(0, r - 1), min(SIZE, r + 2))
                                for cc in range(max(0, c - 1), min(SIZE, c + 2))
                                if (rr, cc) != (r, c))
                if neighbors == 3 or ((r, c) in board and neighbors == 2):
                    nxt.add((r, c))
        board = frozenset(nxt)
        states.append(board)
    return states


def encode(cells):
    return '/'.join(''.join('1' if (r, c) in cells else '0' for c in range(SIZE)) for r in range(SIZE))


def candidate_scores(cells):
    coords = list(cells)
    pairs = list(itertools.combinations(coords, 2))
    adjacency = sum(max(abs(a[0]-b[0]), abs(a[1]-b[1])) == 1 for a,b in pairs)
    spread = sum((a[0]-b[0])**2+(a[1]-b[1])**2 for a,b in pairs)
    # Horizontal concentration rewards arrangements that form rows.
    row_pairs = sum(a[0] == b[0] for a,b in pairs)
    return adjacency, row_pairs, spread


def strategies_for_player(budget, player):
    # Equal, nonoverlapping placement areas, reflected across the horizontal axis.
    # Center row is initially empty; births can subsequently bridge the two areas.
    area = [(r,c) for r in ((0,1) if player == 'A' else (3,4)) for c in range(SIZE)]
    candidates = [frozenset(x) for x in itertools.combinations(area, budget)]
    scores = {x: candidate_scores(x) for x in candidates}
    # Fully specified deterministic selection, with unique patterns per family.
    orderings = {
        'Compact': sorted(candidates, key=lambda x: (-scores[x][0], scores[x][2], encode(x))),
        'Row-oriented': sorted(candidates, key=lambda x: (-scores[x][1], -scores[x][0], encode(x))),
        'Dispersed': sorted(candidates, key=lambda x: (-scores[x][2], scores[x][0], encode(x))),
    }
    chosen = {}
    used = set()
    for family in FAMILIES:
        selected = next(x for x in orderings[family] if x not in used)
        chosen[family] = selected
        used.add(selected)
    return chosen


def pure_equilibria(A, B):
    n, m = A.shape
    return [(i,j) for i in range(n) for j in range(m)
            if A[i,j] >= np.max(A[:,j])-TOL and B[i,j] >= np.max(B[i,:])-TOL]


def feasible_support_payoffs(payoffs, own_support, opp_support):
    """Probability of opponent's strategies making own support best responses."""
    n, m = payoffs.shape
    active = list(own_support)
    opposite = list(opp_support)
    baseline = active[0]
    # variables: opponent mixed probabilities (m), common payoff (1)
    eq_rows = []
    eq_rhs = []
    row = np.zeros(m+1)
    row[:m] = 1
    eq_rows.append(row)
    eq_rhs.append(1)
    for i in active:
        row = np.zeros(m+1)
        row[:m] = payoffs[i,:]
        row[-1] = -1
        eq_rows.append(row)
        eq_rhs.append(0)
    ub_rows = []
    for i in range(n):
        if i not in active:
            row = np.zeros(m+1)
            row[:m] = payoffs[i,:]
            row[-1] = -1
            ub_rows.append(row)
    bounds = [(0,0) if j not in opposite else (1e-8,1) for j in range(m)] + [(None,None)]
    res = linprog(np.zeros(m+1), A_ub=np.array(ub_rows) if ub_rows else None,
                  b_ub=np.zeros(len(ub_rows)) if ub_rows else None,
                  A_eq=np.array(eq_rows), b_eq=np.array(eq_rhs), bounds=bounds, method='highs')
    if not res.success:
        return None
    return np.maximum(0, res.x[:m])


def mixed_equilibria(A, B):
    """Support enumeration for two-player bimatrix games, including unequal supports.

    Degenerate games can have continua of equilibria. We return one feasible
    representative per support pair, not the entire continuum.
    """
    n, m = A.shape
    found = []
    for a_size in range(1,n+1):
        for I in itertools.combinations(range(n), a_size):
            for b_size in range(1,m+1):
                for J in itertools.combinations(range(m), b_size):
                    # A is indifferent among I under q; B is indifferent among J under p.
                    q = feasible_support_payoffs(A, I, J)
                    p = feasible_support_payoffs(B.T, J, I)
                    if p is None or q is None:
                        continue
                    av = A @ q
                    bv = p @ B
                    if max(av) - float(p @ A @ q) > 1e-6:
                        continue
                    if max(bv) - float(p @ B @ q) > 1e-6:
                        continue
                    if any(np.allclose(p, old[0], atol=1e-6) and np.allclose(q, old[1], atol=1e-6)
                           for old in found):
                        continue
                    found.append((p,q))
    return found


def zero_sum_equilibrium(A):
    n,m = A.shape
    # A maximizes guaranteed value v; B minimizes guaranteed value w.
    obj = np.r_[np.zeros(n), -1.]
    aub = np.column_stack((-A.T, np.ones(m)))
    res_a = linprog(obj, A_ub=aub, b_ub=np.zeros(m),
                    A_eq=np.array([[*np.ones(n),0.]]), b_eq=[1.],
                    bounds=[(0,None)]*n+[(None,None)], method='highs')
    obj2 = np.r_[np.zeros(m), 1.]
    bub = np.column_stack((A, -np.ones(n)))
    res_b = linprog(obj2, A_ub=bub, b_ub=np.zeros(n),
                    A_eq=np.array([[*np.ones(m),0.]]), b_eq=[1.],
                    bounds=[(0,None)]*m+[(None,None)], method='highs')
    if not (res_a.success and res_b.success):
        raise RuntimeError('Zero-sum equilibrium solver failed')
    p, q = res_a.x[:n], res_b.x[:m]
    p[p<1e-10] = 0
    q[q<1e-10] = 0
    return p, q, float(p @ A @ q), abs(res_a.x[-1]-res_b.x[-1])


def write_csv(path, rows, headers):
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(rows)


def workbook(output, datasets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in datasets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(headers)
        for row in rows:
            ws.append([str(v) if isinstance(v, (set,frozenset)) else v for v in row])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(name='Aptos', bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.alignment = Alignment(wrap_text=True)
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            width = min(65,max(12,max(len(str(c.value or '')) for c in list(col)[:150])+2))
            ws.column_dimensions[col_letter].width = width
    wb.save(output)


def figures(out, budget, A, B, zero, general):
    fig, axes = plt.subplots(1,2,figsize=(11,4.8),layout='constrained')
    for ax, matrix, title in zip(axes,(A,B),('Survival vs. Extinction: A payoff',
                                           'Mutual preservation: B payoff')):
        img = ax.imshow(matrix,cmap='viridis',aspect='equal')
        ax.set_xticks(range(3), FAMILIES, rotation=22, ha='right')
        ax.set_yticks(range(3), FAMILIES)
        ax.set_xlabel('Player B strategy')
        ax.set_ylabel('Player A strategy')
        ax.set_title(title)
        for i in range(3):
            for j in range(3):
                ax.text(j,i,f'{matrix[i,j]:g}',ha='center',va='center',color='white',fontweight='bold')
        fig.colorbar(img,ax=ax,shrink=.75)
    fig.suptitle(f'Cell budget: {budget} per player | 5×5 finite Conway',fontsize=13)
    fig.savefig(out/f'budget_{budget}_payoff_heatmaps.png',dpi=240,bbox_inches='tight')
    plt.close(fig)

    fig, axes = plt.subplots(1,2,figsize=(10,3.7),layout='constrained')
    for ax, vals, label in zip(axes,(zero[0],zero[1]),('Player A','Player B')):
        ax.bar(FAMILIES,vals)
        ax.set_ylim(0,1.05)
        ax.set_ylabel('Equilibrium probability')
        ax.set_title(label)
        ax.tick_params(axis='x',labelrotation=15)
    fig.suptitle(f'Zero-sum equilibrium mixed strategies — budget {budget}')
    fig.savefig(out/f'budget_{budget}_zero_sum_equilibrium.png',dpi=240,bbox_inches='tight')
    plt.close(fig)


def main():
    parser=argparse.ArgumentParser(description='Two-player Conway game theory study (English).')
    parser.add_argument('--output',default='Conway_Game_Theory_Results',help='Output directory')
    args=parser.parse_args()
    out=Path(args.output).resolve()
    out.mkdir(parents=True,exist_ok=True)
    print('CONWAY GAME THEORY — PHASE 4')
    print('Model 1: survival vs. extinction (zero-sum)')
    print('Model 2: mutual preservation (non-zero-sum counterfactual utility)')
    print('Generation 1 is initial; Generation 10 follows nine B3/S23 transitions.')
    print('WARNING: equilibria apply to selected 3x3 strategy sets only.\n')

    strategy_rows=[]; match_rows=[]; generation_rows=[]; eq_rows=[]; pure_rows=[]; summary=[]
    for budget in BUDGETS:
        A_strats=strategies_for_player(budget,'A')
        B_strats=strategies_for_player(budget,'B')
        for player, strat in (('A',A_strats),('B',B_strats)):
            for family,cells in strat.items():
                strategy_rows.append([budget,player,family,len(cells),
                                      '; '.join(f'({r},{c})' for r,c in sorted(cells)),encode(cells)])
        A=np.zeros((3,3)); B=np.zeros((3,3))
        # Precompute individual baseline outcomes for counterfactual contributions.
        solo_A={f:len(evolve(c)[-1]) for f,c in A_strats.items()}
        solo_B={f:len(evolve(c)[-1]) for f,c in B_strats.items()}
        for i,fa in enumerate(FAMILIES):
            for j,fb in enumerate(FAMILIES):
                a=A_strats[fa]; b=B_strats[fb]
                assert len(a)==len(b)==budget and not a.intersection(b)
                states=evolve(a|b)
                n10=len(states[-1])
                A[i,j]=n10
                # Marginal contribution utility: payoff relative to opponent alone.
                B[i,j]=n10-solo_A[fa]
                ua2=n10-solo_B[fb]
                survival=int(n10>0)
                persistence=sum(bool(s) for s in states)
                match_rows.append([budget,fa,fb,len(a|b),n10,survival,persistence,
                                   n10,-n10,ua2,B[i,j],solo_A[fa],solo_B[fb],
                                   encode(states[0]),encode(states[-1])])
                for gen,state in enumerate(states,1):
                    generation_rows.append([budget,fa,fb,gen,len(state),encode(state)])
        # B is a *different* payoff matrix in the nonzero-sum game.
        nonzero_A=np.array([[next(r[9] for r in match_rows if r[0]==budget and r[1]==fa and r[2]==fb)
                             for fb in FAMILIES] for fa in FAMILIES], dtype=float)
        nonzero_B=B.copy()
        zp,zq,zv,duality=zero_sum_equilibrium(A)
        if duality>1e-6:
            raise AssertionError('Duality gap too large')
        zero_pure=pure_equilibria(A,-A)
        nz_pure=pure_equilibria(nonzero_A,nonzero_B)
        nz_mixed=mixed_equilibria(nonzero_A,nonzero_B)
        for model, eqs in (('Survival_vs_Extinction',zero_pure),('Mutual_Preservation',nz_pure)):
            for i,j in eqs:
                pure_rows.append([budget,model,FAMILIES[i],FAMILIES[j],
                                  float((A if model=='Survival_vs_Extinction' else nonzero_A)[i,j]),
                                  float((-A if model=='Survival_vs_Extinction' else nonzero_B)[i,j])])
        eq_rows.append([budget,'Survival_vs_Extinction','Minimax solution',
                        '; '.join(f'{f}={p:.8f}' for f,p in zip(FAMILIES,zp)),
                        '; '.join(f'{f}={p:.8f}' for f,p in zip(FAMILIES,zq)),
                        round(zv,8),round(-zv,8)])
        for idx,(p,q) in enumerate(nz_mixed,1):
            eq_rows.append([budget,'Mutual_Preservation',f'Support equilibrium {idx}',
                            '; '.join(f'{f}={v:.8f}' for f,v in zip(FAMILIES,p)),
                            '; '.join(f'{f}={v:.8f}' for f,v in zip(FAMILIES,q)),
                            round(float(p @ nonzero_A @ q),8),
                            round(float(p @ nonzero_B @ q),8)])
        summary.append([budget,2*budget,round(2*budget/25,3),
                        round(float(A.mean()),3),int(A.min()),int(A.max()),
                        round(zv,6),len(zero_pure),len(nz_pure),len(nz_mixed),
                        round(duality,10)])
        figures(out,budget,A,nonzero_B,(zp,zq),nz_mixed)
        print(f'Budget {budget}: 9 pairings | zero-sum value={zv:.3f} | '
              f'pure NE={len(zero_pure)} | mutual-preservation pure NE={len(nz_pure)} '
              f'| mutual-preservation representative equilibria={len(nz_mixed)}')

    datasets={
        'Study Summary':(['BudgetPerPlayer','TotalInitialCells','InitialDensity','MeanN10',
                          'MinN10','MaxN10','ZeroSumGameValue','ZeroSumPureNE',
                          'MutualPureNE','MutualRepresentativeNE','ZeroSumDualityGap'],summary),
        'Strategies':(['Budget','Player','Family','CellCount','Coordinates','InitialBoard'],strategy_rows),
        'All Pairings':(['Budget','A_Strategy','B_Strategy','InitialLiveCells','FinalLiveCells',
                         'SurvivedG10','ObservedPersistence','ZeroSumPayoffA','ZeroSumPayoffB',
                         'MutualPayoffA','MutualPayoffB','A_AloneN10','B_AloneN10',
                         'InitialBoard','FinalBoard'],match_rows),
        'Equilibria':(['Budget','Model','Equilibrium','PlayerAProbabilities',
                       'PlayerBProbabilities','ExpectedPayoffA','ExpectedPayoffB'],eq_rows),
        'Pure Nash':(['Budget','Model','PlayerA_Strategy','PlayerB_Strategy','PayoffA','PayoffB'],pure_rows),
        'All Generations':(['Budget','A_Strategy','B_Strategy','Generation','LiveCells','Board'],generation_rows),
        'Methodology':(['Parameter','Definition'],[
            ['Board','5x5 finite; cells outside board dead'],
            ['Rule','Conway B3/S23, synchronous'],
            ['Generation','G1 initial, G10 after nine transitions'],
            ['Budget','3, 4, 5, 6 cells per player'],
            ['Strategies','Three selected families per player and budget; see source selection criteria'],
            ['Placement','A rows 0-1; B rows 3-4; center row initially empty; no overlap'],
            ['Model 1','uA=N10(A union B), uB=-N10(A union B)'],
            ['Model 2','uA=N10(A union B)-N10(B alone); uB=N10(A union B)-N10(A alone)'],
            ['Interpretation','Model 2 is counterfactual marginal contribution, NOT persistent cell ownership'],
            ['Scope','3x3 restricted strategy games; not equilibria in unrestricted 5x5 game'],
            ['Equilibria','Linear programming for zero-sum; support enumeration for non-zero-sum'],
            ['Degeneracy','Support enumeration returns representatives, not full equilibrium continua'],
            ['Density comparison','Phase 3 random configurations are a different sampling design'],
        ])
    }
    for name,(headers,rows) in datasets.items():
        slug=name.lower().replace(' ','_')
        write_csv(out/f'{slug}.csv',rows,headers)
    workbook(out/'conway_game_theory_results.xlsx',datasets)
    metadata={'board_size':SIZE,'observed_generations':OBSERVED_GENERATIONS,
              'transitions':OBSERVED_GENERATIONS-1,'budgets':list(BUDGETS),
              'strategy_families':list(FAMILIES),
              'games_per_budget':9,'total_games':9*len(BUDGETS),
              'zero_sum_utility':'(N10,-N10)',
              'mutual_utility':'(N10(A∪B)-N10(B), N10(A∪B)-N10(A))',
              'source':'conway_game_theory_phase4.py'}
    (out/'experiment_metadata.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
    report=['# Phase 4 — Game-Theoretic Analysis of Conway’s Game of Life','',
            '## Design','',
            'We analyze two explicitly different 3×3 restricted strategy games on a finite 5×5 Conway B3/S23 board. '
            'Generation 1 is the initial configuration and Generation 10 follows nine synchronous updates. '
            'For each budget of 3–6 live cells per player, three deterministic strategy templates '
            '(Compact, Row-oriented, Dispersed) are selected in disjoint upper/lower placement regions. '
            'The empty central row may become occupied through Conway evolution.','',
            '**Model 1 (primary, zero-sum):** Player A maximizes N10; Player B minimizes N10. '
            'Payoffs are (N10, −N10).','',
            '**Model 2 (secondary, non-zero-sum):** Both players maximize their counterfactual '
            'contributions to final population. Payoffs are (N10(A∪B) − N10(B), '
            'N10(A∪B) − N10(A)). These are marginal-contribution utilities, not tracked ownership '
            'of individual cells.','',
            '## Results','',
            '| Cells/player | Total initial | Mean N10 | Zero-sum value | Zero-sum pure NE | Mutual pure NE | Mutual equilibrium representatives |',
            '|---:|---:|---:|---:|---:|---:|---:|']
    for r in summary:
        report.append(f'| {r[0]} | {r[1]} | {r[3]:.3f} | {r[6]:.3f} | {r[7]} | {r[8]} | {r[9]} |')
    report += ['', '## Interpretation and limitations','',
               '- These equilibria are valid **only within the selected three strategies per player**, not all possible 5×5 configurations.',
               '- The spatial partition prevents overlapping initial choices but introduces location and boundary effects. It is a controlled design choice, not a claim of spatial neutrality.',
               '- Strategy-family labels are operational labels selected by deterministic geometric scoring; the patterns are not universal representatives of all clustered/linear/dispersed structures.',
               '- Model 2 does not create persistent ownership, competing cell species, or modify the B3/S23 transition rule.',
               '- The sample includes 36 distinct joint initial configurations, each deterministic. No statistical inference over random strategy samples is justified.',
               '- In degenerate non-zero-sum games, infinitely many equilibria may exist. The report enumerates feasible support representatives only.',
               '- Phase 3 density–persistence results motivate the budgets but do not establish that any budget is optimal for this strategic design.',
               '', '## Reproducibility','',
               'Run `py -m pip install numpy scipy matplotlib openpyxl`, then `py conway_game_theory_phase4.py`.',
               'See `strategies.csv`, `all_pairings.csv`, `all_generations.csv`, `equilibria.csv`, '
               '`pure_nash.csv`, and `experiment_metadata.json` for full auditability.']
    (out/'scientific_report.md').write_text('\n'.join(report)+'\n',encoding='utf-8')
    print(f'\nSUCCESS. Files saved to: {out}')
    print('Excel, 7 CSV files, 8 PNG figures, Markdown report, JSON metadata.')


if __name__=='__main__':
    main()
