#!/usr/bin/env python3
"""Phase 6 — spatial sensitivity of restricted Conway 3x3 games.

Windows:
    py -m pip install numpy scipy matplotlib openpyxl
    py conway_spatial_sensitivity.py
    py conway_spatial_sensitivity.py --samples 100 --seed 20261020

The two players place equal numbers of cells in disjoint 2x5 regions
(rows 0-1 and 3-4). Geometry is changed within each region using
bijective permutations. These are NOT generally symmetries of the full
5x5 Conway board: their effects are precisely what this study measures.
"""
import argparse, csv, json, math, random
from collections import defaultdict
from itertools import combinations, product
from pathlib import Path
import numpy as np
from scipy.optimize import linprog
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

BUDGETS=(3,4,5,6,7)
TRANSFORMS=('identity','horizontal_reflection','column_shift_1','row_swap')
NEIGHBORS=[]
for k in range(25):
    r,c=divmod(k,5)
    NEIGHBORS.append(sum(1<<(rr*5+cc) for rr in range(max(0,r-1),min(5,r+2)) for cc in range(max(0,c-1),min(5,c+2)) if (rr,cc)!=(r,c)))

def evolve(mask):
    for _ in range(9):
        next_mask=0
        for k,neighbors in enumerate(NEIGHBORS):
            n=(mask&neighbors).bit_count()
            if n==3 or (n==2 and ((mask>>k)&1)):
                next_mask|=1<<k
        mask=next_mask
    return mask.bit_count()

def board(mask):
    return '/'.join(''.join('1' if (mask>>(5*r+c))&1 else '0' for c in range(5)) for r in range(5))

def transform(mask,player,kind):
    rows=(0,1) if player=='A' else (3,4)
    result=0
    for r in rows:
        for c in range(5):
            if not ((mask>>(5*r+c))&1): continue
            if kind=='identity':rr,cc=r,c
            elif kind=='horizontal_reflection':rr,cc=r,4-c
            elif kind=='column_shift_1':rr,cc=r,(c+1)%5
            elif kind=='row_swap':rr,cc=rows[1] if r==rows[0] else rows[0],c
            result|=1<<(5*rr+cc)
    assert result.bit_count()==mask.bit_count()
    return result

def minimax(A):
    c=[0.,0.,0.,-1.]
    primal=linprog(c,A_ub=np.column_stack((-A.T,np.ones(3))),b_ub=np.zeros(3),A_eq=[[1,1,1,0]],b_eq=[1],bounds=[(0,None)]*3+[(None,None)],method='highs')
    dual=linprog([0.,0.,0.,1.],A_ub=np.column_stack((A,-np.ones(3))),b_ub=np.zeros(3),A_eq=[[1,1,1,0]],b_eq=[1],bounds=[(0,None)]*3+[(None,None)],method='highs')
    if not primal.success or not dual.success:raise RuntimeError('Linear programming failed')
    p=np.maximum(0,primal.x[:3]);q=np.maximum(0,dual.x[:3]);p/=p.sum();q/=q.sum()
    value=float(p@A@q)
    if abs(primal.x[3]-dual.x[3])>1e-6:raise RuntimeError('Minimax duality gap')
    return p,q,value

def pure_ne(A,B):
    return [(i,j) for i in range(3) for j in range(3) if A[i,j]>=A[:,j].max()-1e-9 and B[i,j]>=B[i,:].max()-1e-9]

def write_csv(path,header,rows):
    with path.open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.writer(f);writer.writerow(header);writer.writerows(rows)

def excel(path,tables):
    wb=Workbook();wb.remove(wb.active)
    for name,(header,rows) in tables.items():
        ws=wb.create_sheet(name[:31]);ws.append(header)
        for row in rows:ws.append([int(x) if isinstance(x,np.integer) else float(x) if isinstance(x,np.floating) else x for x in row])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='183153');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
        for i,col in enumerate(ws.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=min(42,max(13,max(len(str(c.value or '')) for c in list(col)[:80])+2))
    wb.save(path)

def main():
    ap=argparse.ArgumentParser(description='Phase 6: spatial transformations and Nash equilibrium sensitivity')
    ap.add_argument('--samples',type=int,default=40,help='Independent strategy sets per budget (default: 40)')
    ap.add_argument('--seed',type=int,default=20261020)
    ap.add_argument('--output',default='Conway_Spatial_Sensitivity')
    args=ap.parse_args()
    if args.samples<2:ap.error('--samples must be >=2')
    out=Path(args.output).resolve();out.mkdir(parents=True,exist_ok=True)
    rng=random.Random(args.seed)
    regions={'A':[r*5+c for r in (0,1) for c in range(5)],'B':[r*5+c for r in (3,4) for c in range(5)]}
    pools={b:{pl:[sum(1<<k for k in group) for group in combinations(region,b)] for pl,region in regions.items()} for b in BUDGETS}
    cache={}
    def score(mask):
        if mask not in cache:cache[mask]=evolve(mask)
        return cache[mask]
    games=[];pairings=[];strategies=[];paired=[];summaries=[]
    print('PHASE 6 — SPATIAL SENSITIVITY OF NASH EQUILIBRIA')
    print(f'Seed: {args.seed} | Independent sets per budget: {args.samples} | Transform pairs: 16')
    print('Primary: survival versus extinction. Secondary: counterfactual mutual preservation.')
    for budget in BUDGETS:
        per_budget=[];paired_budget=[]
        for sample in range(1,args.samples+1):
            base_a=rng.sample(pools[budget]['A'],3);base_b=rng.sample(pools[budget]['B'],3)
            sample_results={}
            for ta,tb in product(TRANSFORMS,repeat=2):
                pa=[transform(mask,'A',ta) for mask in base_a]
                pb=[transform(mask,'B',tb) for mask in base_b]
                payoff=np.array([[score(a|b) for b in pb] for a in pa],dtype=float)
                p,q,value=minimax(payoff)
                pure=pure_ne(payoff,-payoff)
                solo_a=np.array([score(a) for a in pa]);solo_b=np.array([score(b) for b in pb])
                marginal_a=payoff-solo_b[None,:]
                marginal_b=payoff-solo_a[:,None]
                mutual=pure_ne(marginal_a,marginal_b)
                label=f'{ta}__{tb}'
                row=[budget,sample,ta,tb,round(value,9),len(pure),int(len(pure)==0),len(mutual),round(float(payoff.mean()),5),';'.join(f'{v:.7f}' for v in p),';'.join(f'{v:.7f}' for v in q)]
                games.append(row);per_budget.append(row)
                sample_results[(ta,tb)]=(value,int(len(pure)==0),len(mutual))
                for pl,items,trans in [('A',pa,ta),('B',pb,tb)]:
                    for i,mask in enumerate(items,1):strategies.append([budget,sample,ta,tb,pl,i,mask,board(mask)])
                for i in range(3):
                    for j in range(3):
                        pairings.append([budget,sample,ta,tb,i+1,j+1,int(payoff[i,j]),int(marginal_a[i,j]),int(marginal_b[i,j]),int((i,j) in pure),int((i,j) in mutual)])
            base_v,base_no,base_mutual=sample_results[('identity','identity')]
            vals=[v for v,_,_ in sample_results.values()]
            no_pure=[no for _,no,_ in sample_results.values()]
            paired_row=[budget,sample,round(base_v,9),round(min(vals),9),round(max(vals),9),round(max(vals)-min(vals),9),sum(no_pure),int(len(set(no_pure))>1),base_no,base_mutual]
            paired.append(paired_row);paired_budget.append(paired_row)
        n=len(per_budget)
        vals=np.array([r[4] for r in per_budget]);no=np.array([r[6] for r in per_budget]);changes=np.array([r[7] for r in paired_budget]);ranges=np.array([r[5] for r in paired_budget]);
        summaries.append([budget,args.samples,n,round(float(vals.mean()),6),round(float(vals.std(ddof=1)),6),round(float(no.mean()),6),int(no.sum()),round(float(changes.mean()),6),int(changes.sum()),round(float(ranges.mean()),6),round(float(np.median(ranges)),6)])
        print(f'Budget {budget}: {n} games | no pure NE: {no.mean():.1%} | paired status changes: {changes.mean():.1%} | mean value range: {ranges.mean():.3f}')
    headers={
        'Summary':['CellsPerPlayer','IndependentStrategySets','TransformedGames','MeanGameValue','SDGameValue','FractionNoPureNE','CountNoPureNE','FractionSetsWithNEStatusChange','SetsWithNEStatusChange','MeanWithinSetValueRange','MedianWithinSetValueRange'],
        'Games':['CellsPerPlayer','Sample','TransformA','TransformB','MinimaxValue','PureZeroSumNECount','NoPureZeroSumNE','PureMutualPreservationNECount','MeanPayoff','PlayerAMix','PlayerBMix'],
        'Paired':['CellsPerPlayer','Sample','BaselineGameValue','MinValueAcrossTransforms','MaxValueAcrossTransforms','WithinSetValueRange','TransformsWithoutPureNE','NEStatusChangesAcrossTransforms','BaselineNoPureNE','BaselineMutualPureNECount'],
        'Pairings':['CellsPerPlayer','Sample','TransformA','TransformB','StrategyA','StrategyB','FinalLiveCells','MarginalContributionA','MarginalContributionB','PureZeroSumNE','PureMutualNE'],
        'Strategies':['CellsPerPlayer','Sample','TransformA','TransformB','Player','StrategyIndex','BoardBitmask','BoardRows'],
        'Methods':['Parameter','Value']}
    methods=[['Board','5x5 finite grid; dead cells beyond boundaries'],['Rule','Conway B3/S23'],['Generation convention','Generation 1 initial; Generation 10 after nine transitions'],['Regions','A rows 0-1, B rows 3-4, middle row unused initially'],['Budget','Each player selects the same number of cells from its 10-cell region'],['Strategy sets','Three distinct masks sampled uniformly without replacement for each player'],['Transformations',', '.join(TRANSFORMS)],['Transform design','All 4 x 4 transform pairs for each fixed sampled strategy set'],['Important caveat','Column shift wraps WITHIN a placement region; it is not a Conway-board symmetry'],['Primary payoff','Player A = final live-cell count; Player B = negative of this count'],['Secondary payoff','Each player receives final live-cell count minus outcome from opponent acting alone'],['Equilibrium scope','Pure Nash for both models; minimax mixed equilibrium for zero-sum model'],['Statistical unit','Independent sampled strategy set, not the 16 correlated transformed games'],['Inference','Exploratory; no claim of unrestricted Nash equilibria or causal density threshold'],['Master seed',args.seed],['Independent sets per budget',args.samples]]
    tables={'Summary':(headers['Summary'],summaries),'Games':(headers['Games'],games),'Paired':(headers['Paired'],paired),'Pairings':(headers['Pairings'],pairings),'Strategies':(headers['Strategies'],strategies),'Methods':(headers['Methods'],methods)}
    for name,(header,rows) in tables.items():write_csv(out/f'{name.lower()}.csv',header,rows)
    excel(out/'conway_spatial_sensitivity.xlsx',tables)
    x=[r[0] for r in summaries]
    fig,ax=plt.subplots(figsize=(9,5.5));ax.plot(x,[r[5]*100 for r in summaries],marker='o');ax.set(xticks=x,xlabel='Cells per player',ylabel='Games without pure Nash equilibrium (%)',title='Spatially transformed 3x3 games');ax.grid(alpha=.25);fig.tight_layout();fig.savefig(out/'01_no_pure_ne_by_budget.png',dpi=240);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,5.5));ax.bar(x,[r[7]*100 for r in summaries]);ax.set(xticks=x,ylim=(0,100),xlabel='Cells per player',ylabel='Independent sets with NE-status change (%)',title='Does spatial transformation change pure-equilibrium existence?');fig.tight_layout();fig.savefig(out/'02_ne_status_changes.png',dpi=240);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,5.5));ax.boxplot([[r[5] for r in paired if r[0]==b] for b in x],tick_labels=[str(b) for b in x]);ax.set(xlabel='Cells per player',ylabel='Max minus min minimax value across 16 transforms',title='Within-strategy-set sensitivity to spatial transformations');ax.grid(axis='y',alpha=.25);fig.tight_layout();fig.savefig(out/'03_within_set_value_ranges.png',dpi=240);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,5.5));ax.boxplot([[r[4] for r in games if r[0]==b] for b in x],tick_labels=[str(b) for b in x],showfliers=False);ax.set(xlabel='Cells per player',ylabel='Minimax expected live cells at Generation 10',title='Minimax game values across transformed games');ax.grid(axis='y',alpha=.25);fig.tight_layout();fig.savefig(out/'04_minimax_values.png',dpi=240);plt.close(fig)
    report=['# Phase 6 — Spatial Sensitivity of Restricted Conway Games','',f'- Master seed: {args.seed}',f'- Independent strategy sets per budget: {args.samples}',f'- Budgets: {list(BUDGETS)}','- Transformations: identity, horizontal reflection, cyclic one-column shift, row swap','- Every independent set is evaluated under 16 paired transformations.','- **Important:** 16 transformations of the same strategy set are correlated observations.','', '## Results', '', '| Cells per player | Games | No pure NE | Sets with equilibrium-status changes | Mean within-set value range |','|---:|---:|---:|---:|---:|']
    for r in summaries:report.append(f'| {r[0]} | {r[2]} | {r[5]:.1%} | {r[7]:.1%} | {r[9]:.3f} |')
    report+=['','## Interpretation','', 'Changes in pure-equilibrium status under fixed-budget transformations provide evidence of spatial sensitivity within sampled restricted 3x3 games. These are not equilibria of the full 5x5 strategy space.','', '## Methodological limitations','', '- Placement regions are fixed, disjoint, and separated by an initially empty middle row.','- The transformations are bijections of placement regions, not necessarily symmetries of the complete Conway board.','- Strategy sets are randomly sampled, and no exhaustive equilibrium claim is made.','- Secondary mutual-preservation model reports pure equilibria only; mixed equilibria are not solved.','- The primary zero-sum utility is a modeling choice, not a built-in competitive rule of Conway.','- Budget effects cannot be separated fully from changing reachable pattern geometry without further controls.','', '## Reproducibility','', 'The CSV files contain the full strategy masks, board rows, payoff pairings, and equilibrium summaries. Re-run the script with the same seed and sample count to reproduce the outputs.']
    (out/'scientific_report.md').write_text('\n'.join(report)+'\n',encoding='utf-8')
    (out/'metadata.json').write_text(json.dumps({'seed':args.seed,'samples_per_budget':args.samples,'budgets':BUDGETS,'transformations':TRANSFORMS,'games':len(games),'pairings':len(pairings),'unique_cached_outcomes':len(cache)},indent=2),encoding='utf-8')
    assert len(games)==len(BUDGETS)*args.samples*16
    assert len(pairings)==len(games)*9
    print(f'SUCCESS: {len(games)} transformed games, {len(pairings)} strategy pairings; files saved to {out}')

if __name__=='__main__':main()
