#!/usr/bin/env python3
"""Analyze and audit a 5x5 Conway density experiment.

Searches recursively for Conway_Extended_Study or accepts --input folder/ZIP.
Outputs an English report, an Excel workbook, CSV summaries and publication-ready charts.
Dependencies: openpyxl, matplotlib (pip install openpyxl matplotlib).
"""
import argparse
import csv
import io
import json
import math
import statistics as stats
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

EXPECTED_FILES = ('all_2500_runs.csv', 'all_25000_generation_states.csv', 'density_summary.csv')
BEHAVIORS = ('Extinct', 'Stable', 'Oscillatory', 'Dynamic')


def locate(root):
    root = Path(root).expanduser().resolve()
    if root.is_file() and root.suffix.lower() == '.zip':
        with zipfile.ZipFile(root) as archive:
            names = archive.namelist()
            candidates = [n for n in names if n.endswith('/all_2500_runs.csv')]
            for candidate in candidates:
                prefix = candidate[:-len('all_2500_runs.csv')]
                if all(prefix + f in names for f in EXPECTED_FILES):
                    return ('zip', root, prefix)
        raise FileNotFoundError('ZIP does not contain the required three CSV files.')
    if root.is_file():
        root = root.parent
    folders = [root] + sorted((p for p in root.rglob('all_2500_runs.csv') if p.is_file()), key=str)
    for p in folders:
        folder = p.parent if p.is_file() else p
        if all((folder / f).is_file() for f in EXPECTED_FILES):
            return ('folder', folder, '')
    archives = sorted(root.rglob('Conway_Extended_Study*.zip'))
    for archive in archives:
        try:
            return locate(archive)
        except (FileNotFoundError, zipfile.BadZipFile):
            continue
    raise FileNotFoundError('No valid Conway dataset found. Use --input with the folder or ZIP path.')


def read_bytes(source, name):
    kind, path, prefix = source
    if kind == 'zip':
        with zipfile.ZipFile(path) as archive:
            return archive.read(prefix + name)
    return (path / name).read_bytes()


def load_csv(source, name):
    return list(csv.DictReader(io.StringIO(read_bytes(source, name).decode('utf-8-sig'))))


def board_lines(board):
    lines = board.split('/')
    if len(lines) != 5 or any(len(line) != 5 or set(line) - {'0', '1'} for line in lines):
        raise ValueError(f'Invalid 5x5 board: {board!r}')
    return lines


def evolve(board):
    old = board_lines(board)
    new = []
    for r in range(5):
        row = ''
        for c in range(5):
            count = sum(old[rr][cc] == '1'
                        for rr in range(max(0, r-1), min(5, r+2))
                        for cc in range(max(0, c-1), min(5, c+2))
                        if (rr, cc) != (r, c))
            alive = old[r][c] == '1'
            row += '1' if count == 3 or (alive and count == 2) else '0'
        new.append(row)
    return '/'.join(new)


def live(board):
    return board.count('1')


def classify(states):
    if live(states[-1]) == 0:
        return 'Extinct'
    if states[-1] == states[-2]:
        return 'Stable'
    if states[-1] in states[:-2]:
        return 'Oscillatory'
    return 'Dynamic'


def audit(runs, states, summary):
    issues = []
    by_run = defaultdict(dict)
    for row in states:
        rid = row['RunID']
        gen = int(row['Generation'])
        if gen in by_run[rid]:
            issues.append(f'Duplicate state: {rid} generation {gen}')
        by_run[rid][gen] = row
    seen = set()
    for row in runs:
        rid = row['RunID']
        if rid in seen:
            issues.append(f'Duplicate run ID: {rid}')
        seen.add(rid)
        density = int(row['InitialCells'])
        states_for_run = by_run.get(rid, {})
        if sorted(states_for_run) != list(range(1, 11)):
            issues.append(f'{rid}: expected generations 1..10')
            continue
        boards = [states_for_run[g]['Board'] for g in range(1, 11)]
        try:
            for g, b in enumerate(boards, 1):
                board_lines(b)
                if int(states_for_run[g]['LiveCells']) != live(b):
                    issues.append(f'{rid}: incorrect live-cell count at generation {g}')
                if int(row[f'Live_G{g}']) != live(b):
                    issues.append(f'{rid}: run summary disagrees at generation {g}')
            if live(boards[0]) != density:
                issues.append(f'{rid}: incorrect initial density')
            if boards[0] != row['InitialBoard'] or boards[-1] != row['FinalBoard']:
                issues.append(f'{rid}: initial/final board mismatch')
            for g in range(9):
                if evolve(boards[g]) != boards[g+1]:
                    issues.append(f'{rid}: invalid transition {g+1} to {g+2}')
            persistence = sum(live(b) > 0 for b in boards)
            if int(row['Persistence']) != persistence:
                issues.append(f'{rid}: persistence mismatch')
            if int(row['ReachGeneration10']) != int(live(boards[-1]) > 0):
                issues.append(f'{rid}: survival flag mismatch')
            if int(row['FinalLiveCells']) != live(boards[-1]):
                issues.append(f'{rid}: final live count mismatch')
            if row['Behavior'] != classify(boards):
                issues.append(f'{rid}: behavior mismatch')
        except (ValueError, KeyError) as exc:
            issues.append(f'{rid}: {exc}')
    if len(runs) != 2500: issues.append(f'Expected 2500 runs; found {len(runs)}')
    if len(states) != 25000: issues.append(f'Expected 25000 states; found {len(states)}')
    if len(seen) != len(by_run): issues.append('Run/state ID sets differ')
    by_density = Counter(int(r['InitialCells']) for r in runs)
    for n in range(1, 26):
        if by_density[n] != 100:
            issues.append(f'Density {n}: expected 100 runs, found {by_density[n]}')
    if len(summary) != 25: issues.append('Expected 25 density summary rows')
    return issues


def aggregate(runs, original_summary):
    grouped = defaultdict(list)
    for row in runs:
        grouped[int(row['InitialCells'])].append(row)
    historical = {int(r['InitialCells']): float(r['HistoricalNESMean'])
                  for r in original_summary if r.get('HistoricalNESMean', '').strip()}
    output = []
    for n in sorted(grouped):
        group = grouped[n]
        values = [int(r['Persistence']) for r in group]
        final = [int(r['FinalLiveCells']) for r in group]
        survived = sum(int(r['ReachGeneration10']) for r in group)
        count = len(group)
        mean = stats.mean(values)
        sd = stats.stdev(values) if count > 1 else 0
        margin = 1.96 * sd / math.sqrt(count)
        p = survived / count
        # Wilson 95% binomial confidence interval
        z = 1.96
        denom = 1 + z*z/count
        center = (p + z*z/(2*count)) / denom
        half = z * math.sqrt(p*(1-p)/count + z*z/(4*count*count)) / denom
        behavior = Counter(r['Behavior'] for r in group)
        nes = historical.get(n)
        output.append({
            'InitialCells': n, 'DensityPercent': n*4, 'Runs': count,
            'MeanPersistence': mean, 'SDPersistence': sd,
            'MeanCI95Low': max(0, mean-margin), 'MeanCI95High': min(10, mean+margin),
            'SurvivedGeneration10': survived, 'SurvivalRate': p,
            'SurvivalWilson95Low': max(0, center-half), 'SurvivalWilson95High': min(1, center+half),
            'MeanFinalLiveCells': stats.mean(final),
            **{b: behavior[b] for b in BEHAVIORS},
            'HistoricalNESMean': nes if nes is not None else '',
            'PythonMinusNES': mean-nes if nes is not None else '',
        })
    return output


def save_csv(path, rows):
    with path.open('w', newline='', encoding='utf-8-sig') as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def create_charts(summary, runs, out):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.size': 10, 'savefig.dpi': 300, 'axes.spines.top': False, 'axes.spines.right': False})
    x = [r['InitialCells'] for r in summary]
    y = [r['MeanPersistence'] for r in summary]
    lo = [r['MeanCI95Low'] for r in summary]
    hi = [r['MeanCI95High'] for r in summary]
    historical = [float(r['HistoricalNESMean']) if r['HistoricalNESMean'] != '' else float('nan') for r in summary]
    fig, ax = plt.subplots(figsize=(10, 5.6))
    ax.plot(x, y, 'o-', label='Python (100 runs per density)')
    ax.fill_between(x, lo, hi, alpha=.18, label='Python 95% CI (normal approximation)')
    ax.plot(x, historical, 's--', label='Historical NES (10 runs per density)')
    ax.set(xlabel='Initial live cells (out of 25)', ylabel='Mean observed nonempty generations', title='Persistence by Initial Density: Python vs Historical NES', xticks=x, ylim=(0, 10.5))
    ax.legend(); fig.tight_layout(); fig.savefig(out/'01_persistence_vs_nes.png'); plt.close(fig)

    fig, ax = plt.subplots(figsize=(10, 5.6))
    p = [r['SurvivalRate']*100 for r in summary]
    lower = [r['SurvivalWilson95Low']*100 for r in summary]
    upper = [r['SurvivalWilson95High']*100 for r in summary]
    ax.plot(x, p, 'o-')
    ax.fill_between(x, lower, upper, alpha=.2, label='95% Wilson interval')
    ax.set(xlabel='Initial live cells (out of 25)', ylabel='Survival to generation 10 (%)', title='Survival Probability by Initial Density', xticks=x, ylim=(0, 100))
    ax.legend(); fig.tight_layout(); fig.savefig(out/'02_survival_wilson_intervals.png'); plt.close(fig)

    fig, ax = plt.subplots(figsize=(11, 5.7))
    bottom = [0]*len(x)
    for b in BEHAVIORS:
        values = [r[b]/r['Runs']*100 for r in summary]
        ax.bar(x, values, bottom=bottom, label=b, width=.85)
        bottom = [a+v for a, v in zip(bottom, values)]
    ax.set(xlabel='Initial live cells (out of 25)', ylabel='Runs (%)', title='End-of-Horizon Behavior Classification', xticks=x, ylim=(0, 105))
    ax.legend(ncol=4, loc='upper center', bbox_to_anchor=(.5, -.12)); fig.tight_layout(); fig.savefig(out/'03_behavior_by_density.png'); plt.close(fig)

    fig, ax = plt.subplots(figsize=(11, 5.6))
    groups = defaultdict(list)
    for r in runs: groups[int(r['InitialCells'])].append(int(r['Persistence']))
    ax.boxplot([groups[n] for n in x], positions=x, widths=.65, showfliers=False)
    ax.set(xlabel='Initial live cells (out of 25)', ylabel='Observed nonempty generations (1–10)', title='Distribution of Persistence Across 100 Trials per Density', xticks=x, ylim=(0, 10.7))
    fig.tight_layout(); fig.savefig(out/'04_persistence_boxplots.png'); plt.close(fig)

    # Explicitly descriptive, not a hypothesis test.
    fig, ax = plt.subplots(figsize=(10, 5.6))
    diff = [r['PythonMinusNES'] for r in summary]
    ax.bar(x, [v if v != '' else 0 for v in diff])
    ax.axhline(0, linewidth=1)
    ax.set(xlabel='Initial live cells (out of 25)', ylabel='Python mean minus historical NES mean', title='Descriptive Differences in Mean Persistence', xticks=x)
    fig.tight_layout(); fig.savefig(out/'05_descriptive_python_minus_nes.png'); plt.close(fig)


def create_workbook(summary, issues, runs, out, metadata):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = 'Density Analysis'
    headers = list(summary[0])
    ws.append(headers)
    for row in summary:
        ws.append([row[h] for h in headers])
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.alignment = Alignment(wrap_text=True)
    ws.row_dimensions[1].height = 35
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        ws.column_dimensions[letter].width = min(28, max(14, max(len(str(c.value or '')) for c in list(col)[:100]) + 2))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, float): c.number_format = '0.000'

    sheet = wb.create_sheet('Audit')
    sheet.append(['Check', 'Result'])
    sheet.append(['Runs', len(runs)])
    sheet.append(['Expected runs', 2500])
    sheet.append(['Audit status', 'PASS' if not issues else 'FAIL'])
    sheet.append(['Issue count', len(issues)])
    for issue in issues[:200]: sheet.append(['Issue', issue])

    sheet = wb.create_sheet('Methodology')
    for key, val in metadata.items(): sheet.append([key, str(val)])
    sheet.append(['Analysis', 'Descriptive analysis; no inferential comparison against NES without original trial-level data'])
    sheet.append(['Survival uncertainty', '95% Wilson binomial confidence intervals'])
    sheet.append(['Persistence uncertainty', '95% normal-approximation confidence intervals'])
    sheet.append(['Behavior caveat', 'Dynamic indicates no recurrence detected within the 10 observed states, not proven nonperiodicity'])

    sheet = wb.create_sheet('Highlights')
    peak = max(summary, key=lambda r: r['MeanPersistence'])
    best_survival = max(r['SurvivalRate'] for r in summary)
    for item in [
        ('Metric', 'Value'), ('Total runs', len(runs)),
        ('Peak mean persistence density', peak['InitialCells']),
        ('Peak mean persistence', peak['MeanPersistence']),
        ('Highest observed survival rate', best_survival),
        ('Densities with highest survival', ', '.join(str(r['InitialCells']) for r in summary if r['SurvivalRate'] == best_survival)),
        ('Audit status', 'PASS' if not issues else 'FAIL'),
    ]: sheet.append(item)
    for s in wb:
        if s.title != 'Density Analysis':
            s.column_dimensions['A'].width = 38
            s.column_dimensions['B'].width = 90
            for c in s[1]:
                c.font = Font(bold=True, color='FFFFFF')
                c.fill = PatternFill('solid', fgColor='17365D')
    wb.save(out/'conway_statistical_analysis.xlsx')


def write_report(summary, issues, runs, states, out, source):
    peak = max(summary, key=lambda r: r['MeanPersistence'])
    top_survival = max(r['SurvivalRate'] for r in summary)
    top_densities = [r['InitialCells'] for r in summary if r['SurvivalRate'] == top_survival]
    differences = [abs(r['PythonMinusNES']) for r in summary if r['PythonMinusNES'] != '']
    avg_abs_diff = stats.mean(differences) if differences else None
    text = f'''# Conway's Game of Life: Expanded Density Study — Statistical Review

## Dataset and audit

- Source: `{source[1]}`
- Runs: **{len(runs):,}**; generation records: **{len(states):,}**.
- Audit: **{'PASS' if not issues else 'FAIL'}** ({len(issues)} issue(s)).
- Grid: 5 × 5, finite boundaries, synchronous B3/S23.
- Generation 1 is the initial board; generation 10 follows nine transitions.
- Persistence: number of observed nonempty board states across generations 1–10.

## Main descriptive findings

- Highest observed mean persistence: **{peak['MeanPersistence']:.2f}** at **{peak['InitialCells']} initial cells** ({peak['DensityPercent']}% initial density; 100 trials).
- Approximate 95% confidence interval for that density's mean: **[{peak['MeanCI95Low']:.2f}, {peak['MeanCI95High']:.2f}]**.
- Highest observed generation-10 survival: **{top_survival:.0%}**, at densities **{', '.join(map(str, top_densities))}** cells.
- Overall generation-10 survival: **{sum(int(r['ReachGeneration10']) for r in runs)/len(runs):.1%}**.
- Overall behavior counts: {', '.join(f'{b}: {sum(r[b] for r in summary)}' for b in BEHAVIORS)}.
'''
    if avg_abs_diff is not None:
        text += f'- Mean absolute difference between Python and historical NES density means: **{avg_abs_diff:.2f} generations** (descriptive only).\n'
    text += '''
## Interpretation and limitations

The highest sample mean does not establish a statistically unique optimal density. Neighboring density means may overlap, and comparing the highest of 25 sample means introduces selection effects. The historical NES dataset supplies aggregate means rather than individual run outcomes, so this program does **not** claim a formal between-platform significance test. The comparison also assumes that the original NES persistence definition and sampling method match those used here; confirm this before making stronger claims.

Behavior labels describe the ten-state observation window. A final state classified as Dynamic may enter a stable state or oscillator after generation 10. Extinct means the final state is empty; Stable and Oscillatory are based on observed recurrence.

## Generated files

- `density_analysis.csv`: recalculated per-density statistics.
- `conway_statistical_analysis.xlsx`: summary, audit, methodology, highlights.
- `01_persistence_vs_nes.png`: descriptive NES/Python comparison.
- `02_survival_wilson_intervals.png`: survival and 95% Wilson intervals.
- `03_behavior_by_density.png`: classification proportions.
- `04_persistence_boxplots.png`: trial-level persistence distribution.
- `05_descriptive_python_minus_nes.png`: descriptive mean differences.
- `audit_issues.csv`: detected consistency problems (if any).

## Suggested manuscript wording (English)

A total of 2,500 Python simulations were conducted on a finite 5 × 5 grid, comprising 100 independently sampled initial configurations for each live-cell count from 1 to 25. Ten board states were recorded per simulation, with the initial configuration labeled Generation 1 and the final observation occurring after nine synchronous B3/S23 updates. We summarized persistence, survival to Generation 10, final live-cell population, and end-of-horizon behavior by initial density. Means and approximate 95% confidence intervals were calculated for persistence; Wilson 95% intervals were calculated for survival proportions. Historical NES means were used for descriptive comparison only, because the individual NES trial records were not available for formal cross-platform inference.
'''
    (out/'statistical_report.md').write_text(text, encoding='utf-8')
    (out/'statistical_report.txt').write_text(text, encoding='utf-8')
    with (out/'audit_issues.csv').open('w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f); writer.writerow(['Issue']); writer.writerows([[v] for v in issues])
    return peak, top_survival, top_densities


def main():
    parser = argparse.ArgumentParser(description='Analyze Conway 5x5 simulation outputs and produce an English scientific report.')
    parser.add_argument('--input', default='.', help='Folder containing Conway_Extended_Study, or the study ZIP file (default: current directory).')
    parser.add_argument('--output', default=None, help='Optional output folder. Defaults to Conway_Statistical_Analysis next to the dataset.')
    args = parser.parse_args()
    try:
        source = locate(args.input)
        out = Path(args.output).expanduser().resolve() if args.output else (source[1].parent if source[0] == 'zip' else source[1].parent) / 'Conway_Statistical_Analysis'
        out.mkdir(parents=True, exist_ok=True)
        print(f'Found dataset: {source[1]}')
        print('Loading 2,500 runs and 25,000 generation records...')
        runs = load_csv(source, 'all_2500_runs.csv')
        states = load_csv(source, 'all_25000_generation_states.csv')
        original_summary = load_csv(source, 'density_summary.csv')
        try: metadata = json.loads(read_bytes(source, 'experiment_metadata.json').decode('utf-8-sig'))
        except (FileNotFoundError, KeyError): metadata = {}
        print('Auditing every board, generation transition, and recorded metric...')
        issues = audit(runs, states, original_summary)
        summary = aggregate(runs, original_summary)
        save_csv(out/'density_analysis.csv', summary)
        print('Generating five figures...')
        create_charts(summary, runs, out)
        print('Creating Excel workbook and English scientific report...')
        create_workbook(summary, issues, runs, out, metadata)
        peak, top_survival, top_densities = write_report(summary, issues, runs, states, out, source)
        print('\n=== ANALYSIS COMPLETE ===')
        print(f'Audit: {"PASS" if not issues else "FAIL"} ({len(issues)} issues)')
        print(f'Highest mean persistence: {peak["MeanPersistence"]:.2f} at {peak["InitialCells"]} initial cells')
        print(f'Highest generation-10 survival: {top_survival:.0%} at densities {top_densities}')
        print(f'Output folder: {out}')
        if issues:
            print('WARNING: Check audit_issues.csv before citing these results.')
            return 2
        return 0
    except (FileNotFoundError, zipfile.BadZipFile, ValueError, KeyError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
