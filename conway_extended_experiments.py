#!/usr/bin/env python3
"""Conway's Game of Life: reproducible 5x5 extended density experiment.

Run: py conway_extended_experiments.py
Requires: py -m pip install openpyxl matplotlib

Generation 1 is the initial board; Generation 10 is nine B3/S23 updates.
Boundaries are finite (outside cells are dead). A random board is sampled
uniformly without replacement from all 25 positions for each initial count.
"""
import csv
import hashlib
import json
import math
import random
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean, stdev
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

GRID_SIZE = 5
GENERATIONS = 10  # generation 1 is initial; generation 10 = 9 updates
REPLICATES = 100
MASTER_SEED = 20261020
OUTPUT_ROOT = Path(__file__).resolve().parent / 'Conway_Extended_Study'
# Historic NES density means (10 trials per count), preserved as previously reported.
# This is a reference series, NOT a reconstruction of the original raw NES trials.
NES_MEANS = [1.0,1.0,1.9,2.7,4.7,7.0,7.2,7.8,8.3,8.6,9.3,8.0,9.1,9.0,8.3,6.9,6.5,5.4,4.2,3.1,2.3,2.0,2.0,2.0,2.0]


def update(board):
    new = []
    for r in range(GRID_SIZE):
        row = []
        for c in range(GRID_SIZE):
            neighbors = sum(board[rr][cc] for rr in range(max(0,r-1), min(GRID_SIZE,r+2))
                            for cc in range(max(0,c-1), min(GRID_SIZE,c+2)) if (rr,cc)!=(r,c))
            row.append(int(neighbors == 3 or (board[r][c] == 1 and neighbors == 2)))
        new.append(tuple(row))
    return tuple(new)


def random_board(n, seed):
    positions = random.Random(seed).sample(range(25), n)
    chosen = set(positions)
    return tuple(tuple(int(5*r+c in chosen) for c in range(5)) for r in range(5))


def serialize(board):
    return '/'.join(''.join(map(str,row)) for row in board)


def classification(states):
    """Classify the observed trajectory; categories are operational definitions.

    Extinct: zero live cells in the final state.
    Stable: nonempty final state equals preceding state.
    Oscillatory: nonempty final state repeats an earlier state (period >=2).
    Dynamic: nonempty final state with no detected recurrence at the end.
    """
    final = states[-1]
    if not any(sum(row) for row in final): return 'Extinct'
    if final == states[-2]: return 'Stable'
    if final in states[:-2]: return 'Oscillatory'
    return 'Dynamic'


def persistence(states):
    """Count observed nonempty states from generations 1..10.
    Conway extinction is absorbing, so this is min(first empty index,10).
    """
    return next((i for i,board in enumerate(states) if not any(map(sum,board))), len(states))


def write_csv(path, headers, rows):
    with path.open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.writer(f); writer.writerow(headers); writer.writerows(rows)


def ci95(values):
    if len(values)<2: return (mean(values),mean(values))
    # Normal approximation to mean CI, documented and used descriptively.
    delta=1.96*stdev(values)/math.sqrt(len(values))
    return mean(values)-delta,mean(values)+delta


def plot_results(folder, summaries, runs):
    x=[r['InitialCells'] for r in summaries]
    py=[r['MeanPersistence'] for r in summaries]
    lo=[r['CI95Low'] for r in summaries]
    hi=[r['CI95High'] for r in summaries]
    fig,ax=plt.subplots(figsize=(10,5.5))
    ax.plot(x,py,marker='o',label='Python (100 trials per density)')
    ax.fill_between(x,lo,hi,alpha=.2,label='Python 95% CI (normal approximation)')
    ax.plot(x,NES_MEANS,marker='s',linestyle='--',alpha=.8,label='Historical NES mean (10 trials per density)')
    ax.set(xlabel='Initial live cells (out of 25)',ylabel='Persistence (observed nonempty generations)',title='Persistence by initial density: Python and historical NES')
    ax.set_xticks(x);ax.set_ylim(0,10.5);ax.grid(alpha=.25);ax.legend();fig.tight_layout()
    fig.savefig(folder/'01_persistence_comparison.png',dpi=220);plt.close(fig)
    fig,ax=plt.subplots(figsize=(10,5.5))
    ax.plot(x,[100*r['ReachGeneration10Rate'] for r in summaries],marker='o')
    ax.set(xlabel='Initial live cells (out of 25)',ylabel='Runs alive at Generation 10 (%)',title='Generation 10 survival rate (100 runs per density)')
    ax.set_xticks(x);ax.set_ylim(0,105);ax.grid(alpha=.25);fig.tight_layout()
    fig.savefig(folder/'02_survival_rate.png',dpi=220);plt.close(fig)
    cats=['Extinct','Stable','Oscillatory','Dynamic']
    fig,ax=plt.subplots(figsize=(10,5.5));bottom=[0]*25
    for cat in cats:
        counts=[r[cat] for r in summaries]
        ax.bar(x,counts,bottom=bottom,label=cat)
        bottom=[a+b for a,b in zip(bottom,counts)]
    ax.set(xlabel='Initial live cells (out of 25)',ylabel='Runs (out of 100)',title='Observed final behavior by initial density')
    ax.set_xticks(x);ax.legend();fig.tight_layout()
    fig.savefig(folder/'03_behavior_distribution.png',dpi=220);plt.close(fig)
    fig,ax=plt.subplots(figsize=(10,5.5))
    by_n={n:[] for n in x}
    for row in runs: by_n[row['InitialCells']].append(row['FinalLiveCells'])
    ax.boxplot([by_n[n] for n in x],positions=x,showfliers=False)
    ax.set(xlabel='Initial live cells (out of 25)',ylabel='Live cells at Generation 10',title='Distribution of final live-cell counts')
    ax.set_xticks(x);ax.grid(axis='y',alpha=.25);fig.tight_layout()
    fig.savefig(folder/'04_final_population_boxplot.png',dpi=220);plt.close(fig)


def write_excel(path, run_rows, generation_rows, summary_rows, metadata):
    """Create an Excel workbook with openpyxl (available on PyPI)."""
    wb = Workbook()
    wb.remove(wb.active)

    def add(name, headers, rows, widths=None):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30
        for cell in ws[1]:
            cell.fill = PatternFill('solid', fgColor='12314A')
            cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        if widths:
            for index, width in enumerate(widths, start=1):
                ws.column_dimensions[get_column_letter(index)].width = width
        return ws

    summary_headers = ['InitialCells','DensityPct','Runs','MeanPersistence',
        'SDPersistence','CI95Low','CI95High','ReachGeneration10Rate',
        'MeanFinalLiveCells','Extinct','Stable','Oscillatory','Dynamic',
        'HistoricalNESMean','PythonMinusNES']
    summary_rows = [[r[h] for h in summary_headers] for r in summary_rows]
    summary = add('Density Summary', summary_headers, summary_rows,
        [15,14,9,19,17,14,14,24,22,12,12,16,12,20,19])
    for i in range(2, len(summary_rows) + 2):
        summary[f'O{i}'] = f'=D{i}-N{i}'
        for col in ('B','H'):
            summary[f'{col}{i}'].number_format = '0.0%'
        for col in ('D','E','F','G','I','N','O'):
            summary[f'{col}{i}'].number_format = '0.00'

    run_headers = ['RunID','InitialCells','DensityPct','Trial','Seed',
        'InitialBoard','FinalBoard','Persistence','ReachGeneration10',
        'FinalLiveCells','Behavior'] + [f'Live_G{i}' for i in range(1,11)]
    all_runs = add('All Runs', run_headers, run_rows,
        [14,14,14,10,24,35,35,14,20,16,17] + [11]*10)
    for i in range(2, len(run_rows)+2):
        all_runs[f'C{i}'].number_format = '0.0%'

    add('Methodology', ['Parameter','Value'],
        [[k, str(v)] for k,v in metadata.items()], [35,95])
    add('Historical NES', ['InitialCells','NESMeanPersistence','NESRuns','SourceNote'],
        [[i,NES_MEANS[i-1],10,
          'Previously reported aggregate; original raw NES runs not included']
          for i in range(1,26)], [18,25,15,75])
    # The complete 25,000-state trajectory is exported to CSV, avoiding
    # a large redundant worksheet. Excel contains all 2,500 run summaries.
    wb.save(path)


def main():
    OUTPUT_ROOT.mkdir(parents=True,exist_ok=True)
    run_rows=[];gen_rows=[];runs=[];summaries=[]
    for n in range(1,26):
        density=[]
        for trial in range(1,REPLICATES+1):
            # Stable unique seed per (n,trial), independent of execution order.
            seed=int.from_bytes(hashlib.sha256(f'{MASTER_SEED}:{n}:{trial}'.encode()).digest()[:8],'big')
            board=random_board(n,seed)
            states=[board]
            for _ in range(GENERATIONS-1): states.append(update(states[-1]))
            counts=[sum(map(sum,s)) for s in states]
            p=persistence(states); behavior=classification(states)
            run_id=f'N{n:02d}_R{trial:03d}'
            record=dict(RunID=run_id,InitialCells=n,DensityPct=n/25,Trial=trial,Seed=str(seed),
                        InitialBoard=serialize(states[0]),FinalBoard=serialize(states[-1]),
                        Persistence=p,ReachGeneration10=int(counts[-1]>0),FinalLiveCells=counts[-1],Behavior=behavior)
            runs.append(record);density.append(record)
            run_rows.append([run_id,n,n/25,trial,str(seed),record['InitialBoard'],record['FinalBoard'],p,record['ReachGeneration10'],counts[-1],behavior,*counts])
            for g,s in enumerate(states,1): gen_rows.append([run_id,n,trial,g,serialize(s),counts[g-1]])
        pvals=[r['Persistence'] for r in density]
        lo,hi=ci95(pvals)
        hist=Counter(r['Behavior'] for r in density)
        summaries.append(dict(InitialCells=n,DensityPct=n/25,Runs=len(density),MeanPersistence=mean(pvals),
                              SDPersistence=stdev(pvals),CI95Low=max(0,lo),CI95High=min(10,hi),
                              ReachGeneration10Rate=mean(r['ReachGeneration10'] for r in density),
                              MeanFinalLiveCells=mean(r['FinalLiveCells'] for r in density),
                              **{cat:hist[cat] for cat in ('Extinct','Stable','Oscillatory','Dynamic')},
                              HistoricalNESMean=NES_MEANS[n-1],PythonMinusNES=mean(pvals)-NES_MEANS[n-1]))
        print(f'Density {n:2d}/25 | 100 trials | mean persistence {summaries[-1]["MeanPersistence"]:.2f} | alive at G10 {summaries[-1]["ReachGeneration10Rate"]:.0%}')
    run_headers=['RunID','InitialCells','DensityPct','Trial','Seed','InitialBoard','FinalBoard','Persistence','ReachGeneration10','FinalLiveCells','Behavior']+[f'Live_G{i}' for i in range(1,11)]
    gen_headers=['RunID','InitialCells','Trial','Generation','Board','LiveCells']
    summary_headers=list(summaries[0].keys())
    write_csv(OUTPUT_ROOT/'all_2500_runs.csv',run_headers,run_rows)
    write_csv(OUTPUT_ROOT/'all_25000_generation_states.csv',gen_headers,gen_rows)
    write_csv(OUTPUT_ROOT/'density_summary.csv',summary_headers,[[r[h] for h in summary_headers] for r in summaries])
    metadata={'Experiment':'Expanded Conway density experiment','Grid':'5x5','Boundary':'Finite, outside cells dead',
              'TransitionRule':'B3/S23, synchronous','InitialStateLabel':'Generation 1','LastStateLabel':'Generation 10',
              'UpdatesPerRun':GENERATIONS-1,'InitialCells':'1..25','ReplicatesPerDensity':REPLICATES,
              'TotalRuns':25*REPLICATES,'MasterSeed':MASTER_SEED,
              'Sampling':'Uniform random positions without replacement, exactly n live cells',
              'Persistence':'Number of nonempty observed states (1..10); 0 if initially empty',
              'ReachGeneration10':'1 iff final board is nonempty',
              'Stable':'Nonempty final board equals preceding board',
              'Oscillatory':'Nonempty final board equals earlier board, excluding preceding board',
              'Dynamic':'Nonempty final board has no detected recurrence',
              'Extinct':'Final board empty',
              'CI95':'Mean +/- 1.96 * sample SD / sqrt(100); normal approximation',
              'NESReference':'Historical aggregate means (10 trials per density); not raw data',
              'ComparabilityCaution':'Verify NES persistence convention and random sampling procedure before inferential comparisons',
              'GeneratedUTC':datetime.now(timezone.utc).isoformat()}
    (OUTPUT_ROOT/'experiment_metadata.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
    write_excel(OUTPUT_ROOT/'conway_extended_results.xlsx',run_rows,gen_rows,summaries,metadata)
    plot_results(OUTPUT_ROOT,summaries,runs)
    print('\nSUCCESS: 2,500 runs, 25,000 board states, Excel, 3 CSV files, 4 PNG charts.')
    print('Output directory:',OUTPUT_ROOT)

if __name__=='__main__': main()
