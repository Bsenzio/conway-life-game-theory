#!/usr/bin/env python3
"""Phase 5: reproducible robustness study of restricted 3x3 Conway games.

Windows: py -m pip install numpy scipy matplotlib openpyxl
         py conway_robustness_analysis.py
Options: py conway_robustness_analysis.py --samples 100 --seed 20261020

This script does NOT claim equilibrium in the unrestricted 5x5 game.
"""
import argparse, csv, json, math, random
from pathlib import Path
from collections import Counter
import numpy as np
from scipy.optimize import linprog
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SIZE=5
GENERATIONS=10
BUDGETS=(3,4,5,6,7)
NAMES=('Strategy 1','Strategy 2','Strategy 3')

# Bitmask representation, with 25 cells indexed row-major.
NEIGHBORS=[]
for k in range(25):
    r,c=divmod(k,5)
    NEIGHBORS.append(sum(1<<(rr*5+cc) for rr in range(max(0,r-1),min(5,r+2))
                         for cc in range(max(0,c-1),min(5,c+2)) if (rr,cc)!=(r,c)))

def evolve(mask):
    x=mask
    for _ in range(GENERATIONS-1):
        y=0
        for k,n in enumerate(NEIGHBORS):
            count=(x & n).bit_count()
            if count==3 or (count==2 and (x>>k)&1):
                y|=1<<k
        x=y
    return x.bit_count()

def board(mask):
    return '/'.join(''.join('1' if mask>>(r*5+c)&1 else '0' for c in range(5)) for r in range(5))

def pure_ne(A,B):
    return [(i,j) for i in range(3) for j in range(3)
            if A[i,j]>=max(A[:,j])-1e-8 and B[i,j]>=max(B[i,:])-1e-8]

def minimax(A):
    # Maximize v s.t. A^T p >= v, sum(p)=1.
    c=np.array([0.,0.,0.,-1.])
    u=np.column_stack((-A.T,np.ones(3)))
    x=linprog(c,A_ub=u,b_ub=np.zeros(3),A_eq=[[1,1,1,0]],b_eq=[1],
              bounds=[(0,None)]*3+[(None,None)],method='highs')
    # Minimize w s.t. A q <= w, sum(q)=1.
    y=linprog([0,0,0,1],A_ub=np.column_stack((A,-np.ones(3))),b_ub=np.zeros(3),
              A_eq=[[1,1,1,0]],b_eq=[1],bounds=[(0,None)]*3+[(None,None)],method='highs')
    if not x.success or not y.success: raise RuntimeError('Minimax linear program failed')
    p=np.maximum(x.x[:3],0);q=np.maximum(y.x[:3],0)
    p/=p.sum();q/=q.sum()
    gap=abs(x.x[3]-y.x[3]);assert gap<1e-6
    return p,q,float(p@A@q),gap

def ci_mean(vals):
    a=np.asarray(vals,dtype=float)
    mean=float(a.mean())
    # Normal-approximation CI for the sample mean; not a population-free guarantee.
    half=1.96*float(a.std(ddof=1))/math.sqrt(len(a)) if len(a)>1 else 0.
    return mean,mean-half,mean+half

def wilson(success,n):
    z=1.96;p=success/n;den=1+z*z/n
    mid=(p+z*z/(2*n))/den
    h=z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/den
    return mid-h,mid+h

def write_csv(path,headers,rows):
    with path.open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f);w.writerow(headers);w.writerows(rows)

def make_excel(path,tables):
    wb=Workbook();wb.remove(wb.active)
    for title,(header,rows) in tables.items():
        ws=wb.create_sheet(title[:31]);ws.append(header)
        for row in rows: ws.append(list(row))
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=PatternFill('solid',fgColor='17365D')
            cell.font=Font(name='Aptos',bold=True,color='FFFFFF')
            cell.alignment=Alignment(wrap_text=True)
        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width=min(55,max(12,max(len(str(c.value or '')) for c in list(col)[:100])+2))
    wb.save(path)

def plot_results(out,summary,records):
    budgets=[r[0] for r in summary]
    def plot_line(idx,label,filename,ylab,percent=False):
        fig,ax=plt.subplots(figsize=(8.5,5.1))
        ax.plot(budgets,[r[idx] for r in summary],marker='o',linewidth=2)
        ax.set_xticks(budgets);ax.set_xlabel('Cells per player')
        ax.set_ylabel(ylab);ax.set_title(label);ax.grid(alpha=.25)
        if percent:ax.set_ylim(0,1)
        fig.tight_layout();fig.savefig(out/filename,dpi=240);plt.close(fig)
    plot_line(3,'Fraction of sampled games without a pure zero-sum Nash equilibrium',
              '01_no_pure_equilibrium_rate.png','Fraction of sampled 3x3 games',True)
    plot_line(6,'Mean minimax value by cell budget','02_mean_game_value.png',
              'Expected live cells at Generation 10')
    plot_line(9,'Mean counterfactual-preservation pure-equilibrium count',
              '03_mutual_preservation_pure_ne.png','Number of pure equilibria per game')
    fig,ax=plt.subplots(figsize=(9,5.5))
    values=[[r['Value'] for r in records if r['Budget']==b] for b in budgets]
    ax.boxplot(values,tick_labels=[str(b) for b in budgets],showfliers=False)
    ax.set(xlabel='Cells per player',ylabel='Minimax game value (final live cells)',
           title='Distribution of zero-sum minimax values across sampled strategy sets')
    ax.grid(axis='y',alpha=.25);fig.tight_layout()
    fig.savefig(out/'04_game_value_boxplots.png',dpi=240);plt.close(fig)
    fig,ax=plt.subplots(figsize=(9,5.5))
    for b in budgets:
        subset=[r for r in records if r['Budget']==b]
        a=sum(r['ZeroSumPureNE']==0 for r in subset);lo,hi=wilson(a,len(subset))
        ax.errorbar(b,a/len(subset),yerr=[[a/len(subset)-lo],[hi-a/len(subset)]],
                    fmt='o',capsize=5)
    ax.set(xticks=budgets,ylim=(0,1),xlabel='Cells per player',
           ylabel='Fraction without pure NE (95% Wilson CI)',
           title='Sampled-game proportion and uncertainty')
    ax.grid(alpha=.25);fig.tight_layout()
    fig.savefig(out/'05_no_pure_ne_wilson.png',dpi=240);plt.close(fig)

def main():
    ap=argparse.ArgumentParser(description='Conway Phase 5 robustness and sensitivity study')
    ap.add_argument('--samples',type=int,default=100,help='Independent sampled 3x3 games per budget')
    ap.add_argument('--seed',type=int,default=20261020,help='Reproducible master seed')
    ap.add_argument('--output',default='Conway_Robustness_Analysis',help='Output folder')
    args=ap.parse_args()
    if args.samples<2:ap.error('--samples must be at least 2')
    out=Path(args.output).resolve();out.mkdir(parents=True,exist_ok=True)
    rng=random.Random(args.seed)
    regions={'A':[r*5+c for r in (0,1) for c in range(5)],
             'B':[r*5+c for r in (3,4) for c in range(5)]}
    from itertools import combinations
    pools={b:{player:[sum(1<<k for k in group) for group in combinations(indices,b)]
              for player,indices in regions.items()} for b in BUDGETS}
    cache={};hits=0
    def outcome(mask):
        nonlocal hits
        if mask not in cache:cache[mask]=evolve(mask)
        else:hits+=1
        return cache[mask]
    record_rows=[];strategy_rows=[];pair_rows=[];equilibrium_rows=[];summary=[];records=[]
    print('PHASE 5 — ROBUSTNESS AND SENSITIVITY ANALYSIS')
    print('Primary: survival versus extinction (zero-sum)')
    print('Secondary: mutual preservation (counterfactual marginal contribution)')
    print(f'Master seed: {args.seed} | Samples per budget: {args.samples}')
    print('Sampling: 3 distinct strategies per player from 2x5 nonoverlapping regions')
    for b in BUDGETS:
        group=[]
        for sample in range(1,args.samples+1):
            # Randomly sampled restricted strategy sets; no rejection based on outcome.
            pa=rng.sample(pools[b]['A'],3);pb=rng.sample(pools[b]['B'],3)
            solo_a=[outcome(a) for a in pa];solo_b=[outcome(z) for z in pb]
            A=np.array([[outcome(a|z) for z in pb] for a in pa],dtype=float)
            MA=A-np.array(solo_b)[None,:]
            MB=A-np.array(solo_a)[:,None]
            p,q,value,gap=minimax(A)
            pure=pure_ne(A,-A);secondary=pure_ne(MA,MB)
            pure_count=len(pure);secondary_count=len(secondary)
            record={'Budget':b,'Sample':sample,'Value':value,
                    'ZeroSumPureNE':pure_count,'MutualPureNE':secondary_count}
            records.append(record);group.append(record)
            record_rows.append([b,sample,round(value,9),pure_count,int(pure_count==0),
                                secondary_count,round(float(A.mean()),5),int(A.min()),int(A.max()),
                                ';'.join(f'{v:.9f}' for v in p),';'.join(f'{v:.9f}' for v in q),gap])
            equilibrium_rows.append([b,sample,'Survival_vs_Extinction',
                ';'.join(f'{v:.9f}' for v in p),';'.join(f'{v:.9f}' for v in q),round(value,9),
                'Minimax LP; valid only within sampled 3x3 strategy set'])
            for player,patterns in [('A',pa),('B',pb)]:
                for idx,mask in enumerate(patterns,1):
                    strategy_rows.append([b,sample,player,idx,mask,board(mask),outcome(mask)])
            for i in range(3):
                for j in range(3):
                    pair_rows.append([b,sample,i+1,j+1,board(pa[i]|pb[j]),int(A[i,j]),
                                      int(MA[i,j]),int(MB[i,j]),int((i,j) in pure),
                                      int((i,j) in secondary)])
        no_pure=sum(x['ZeroSumPureNE']==0 for x in group)
        values=[x['Value'] for x in group]
        mean,lo,hi=ci_mean(values)
        wlo,whi=wilson(no_pure,args.samples)
        secondary_mean=sum(x['MutualPureNE'] for x in group)/args.samples
        summary.append([b,args.samples,no_pure,no_pure/args.samples,wlo,whi,
                        mean,lo,hi,secondary_mean,
                        sum(x['ZeroSumPureNE'] for x in group)/args.samples,
                        sum(x['Value']>1e-8 for x in group)/args.samples])
        print(f'Budget {b}: no pure NE={no_pure}/{args.samples} ({no_pure/args.samples:.1%}); '
              f'mean minimax={mean:.3f}; mutual pure NE mean={secondary_mean:.2f}')
    tables={
      'Summary':(['Budget','SampledGames','NoPureNEGames','NoPureNERate','Wilson95Lower','Wilson95Upper',
                  'MeanGameValue','MeanValue95Lower','MeanValue95Upper','MeanMutualPureNE',
                  'MeanZeroSumPureNE','PositiveValueFraction'],summary),
      'All Games':(['Budget','Sample','ZeroSumGameValue','ZeroSumPureNE','NoPureNE','MutualPureNE',
                    'MeanN10','MinN10','MaxN10','AProbabilities','BProbabilities','DualityGap'],record_rows),
      'Strategies':(['Budget','Sample','Player','StrategyIndex','Bitmask','InitialBoard','AloneN10'],strategy_rows),
      'Payoff Pairings':(['Budget','Sample','AIndex','BIndex','InitialBoard','N10','MutualPayoffA',
                          'MutualPayoffB','ZeroSumPureNE','MutualPureNE'],pair_rows),
      'Equilibria':(['Budget','Sample','Model','AProbabilities','BProbabilities','ExpectedPayoffA','Scope'],equilibrium_rows),
      'Methodology':(['Parameter','Definition'],[
          ['Game','Finite 3x3 sampled strategy sets; zero-sum and counterfactual mutual preservation'],
          ['Board','5x5, finite boundaries; Conway B3/S23'],
          ['Horizon','Generation 1 initial; Generation 10 after 9 synchronous transitions'],
          ['Budgets','3,4,5,6,7 cells per player'],
          ['Regions','Player A rows 0-1; Player B rows 3-4; central row initially empty'],
          ['Sampling','Uniform random 3 distinct patterns per player from all budget-compatible patterns in own region'],
          ['Samples per budget',str(args.samples)],['Master seed',str(args.seed)],
          ['Primary payoffs','A=N10(A union B), B=-N10(A union B)'],
          ['Secondary payoffs','A=N10(A union B)-N10(B alone); B=N10(A union B)-N10(A alone)'],
          ['Primary equilibrium','Zero-sum minimax linear programming; pure NE counted separately'],
          ['Secondary equilibrium','Pure equilibria counted; mixed nonzero-sum equilibria NOT enumerated'],
          ['Confidence intervals','Wilson 95% for sampled no-pure proportion; normal 95% for mean game value'],
          ['Limitation','Restricted sampled 3x3 games; NOT exhaustive Nash equilibria of unrestricted Conway'],
          ['Limitation','Different samples can reuse patterns; each sampled game is an independent random draw'],
          ['Limitation','Spatial partitions and sampling frame affect outcomes'],
      ])}
    for name,(headers,rows) in tables.items():
        write_csv(out/(name.lower().replace(' ','_')+'.csv'),headers,rows)
    make_excel(out/'conway_robustness_results.xlsx',tables)
    plot_results(out,summary,records)
    metadata={'master_seed':args.seed,'samples_per_budget':args.samples,'budgets':BUDGETS,
              'total_sampled_games':len(records),'total_pairings':len(pair_rows),
              'generations':GENERATIONS,'transitions':GENERATIONS-1,
              'strategy_space':'3 strategies/player sampled from 2x5 disjoint zones',
              'mixed_secondary_equilibria':'not enumerated'}
    (out/'experiment_metadata.json').write_text(json.dumps(metadata,indent=2),encoding='utf-8')
    report=['# Phase 5: Nash Equilibrium Robustness and Sensitivity Analysis','',
            '## Methods','',
            f'{args.samples} independently sampled restricted 3x3 games were evaluated at each '
            'cell budget (3–7 cells per player). Player A selected cells in the top two rows '
            'and Player B in the bottom two rows, leaving the central row initially empty. '
            'Within each budget and player region, three distinct initial patterns were '
            'sampled uniformly without replacement per game. All sampling was reproducible '
            f'using master seed {args.seed}. Conway B3/S23 evolved for nine transitions '
            '(Generation 1 initial; Generation 10 final).','',
            'The primary zero-sum model assigns payoffs (N10, -N10), with minimax mixed '
            'strategies computed by linear programming. The secondary model assigns '
            'counterfactual marginal-contribution payoffs; pure equilibria were counted, '
            'but secondary mixed equilibria were not enumerated.','',
            '## Results','',
            '| Cells/player | Games | No pure zero-sum NE | Fraction (Wilson 95% CI) | Mean minimax value (95% CI) | Mean secondary pure NE |',
            '|---:|---:|---:|---:|---:|---:|']
    for r in summary:
        report.append(f'| {r[0]} | {r[1]} | {r[2]} | {r[3]:.3f} [{r[4]:.3f}, {r[5]:.3f}] | '
                      f'{r[6]:.3f} [{r[7]:.3f}, {r[8]:.3f}] | {r[9]:.3f} |')
    report += ['', '## Interpretation','',
       'The analysis quantifies how often the absence of a pure-strategy Nash equilibrium '
       'occurs across randomly sampled *restricted* 3x3 strategy sets. A mixed minimax '
       'equilibrium exists for every zero-sum finite game, even when pure equilibria exist. '
       'Differences across budgets are descriptive and conditional on the specified '
       'sampling and spatial-partition protocol.','',
       '## Limitations','',
       '- The sampled strategy sets are not exhaustive. Results cannot establish a unique critical budget or unrestricted-game equilibrium.',
       '- Each game has three choices per player; the three are sampled configurations, not fixed geometric families.',
       '- The top/bottom partition avoids overlap but may create boundary asymmetry.',
       '- The secondary model does not track persistent cell ownership, and only its pure Nash equilibria are evaluated.',
       '- Confidence intervals describe variation under the chosen random sampling procedure, not uncertainty about deterministic Conway rules.',
       '- The 2,500 density experiments in Phase 3 used a different sampling frame; direct causal comparison is not warranted.',
       '', '## Reproducibility','',
       f'Run: `py conway_robustness_analysis.py --samples {args.samples} --seed {args.seed}`',
       'Outputs include raw strategy boards, all payoff pairings, minimax probabilities, '
       'an Excel workbook, CSV files, and five figures.']
    (out/'scientific_report.md').write_text('\n'.join(report)+'\n',encoding='utf-8')
    print(f'COMPLETE: {len(records)} games, {len(pair_rows)} payoff pairings')
    print(f'Output folder: {out}')
    print(f'Cached distinct Conway simulations: {len(cache)}')

if __name__=='__main__':main()
